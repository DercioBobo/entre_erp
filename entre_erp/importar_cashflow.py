"""One-off import of the old cashflow.xlsx into Plano de Pagamentos.

    bench --site your-site.com execute entre_erp.importar_cashflow.importar \\
        --kwargs "{'caminho': '/caminho/para/cashflow.xlsx', 'ano': 2026}"

Each month sheet (Janeiro … Outubro) becomes one Plano de Pagamentos and the
Investimentos sheet becomes the year's Investimentos plan. Existing plans are
skipped unless `substituir=True`. The month comes from the sheet name, not
cell B1 (B1 is swapped on the Setembro/Outubro sheets).
"""

import re
import unicodedata

import frappe
from frappe.utils import getdate

from entre_erp.install import create_categorias_de_despesa
from entre_erp.pagamentos import MESES, nome_plano, numero_mes

METODOS_DE_PAGAMENTO = ["STB", "BIM", "BIM C"]

ESTADOS = {
	"pago": "Pago",
	"pendente": "Pendente",
	"reservado": "Reservado",
	"parcialmente pago": "Parcialmente Pago",
	"em progresso": "Em Progresso",
	"em espera": "Em Espera",
	"proximo mes": "Próximo Mês",
}

# (name, categoria, valor_padrao, tipo_valor, prioridade, metodo, referencia, pattern)
# Values are the latest ones in the spreadsheet (Setembro/Outubro 2026).
DESPESAS_RECORRENTES = [
	("INSS-DH", "Impostos", 3384.15, "Fixo", "1", None, None, r"^inss\s*-\s*dh"),
	("IRPS-DH", "Impostos", 4844, "Fixo", "1", None, None, r"^irps\s*-\s*dh"),
	("INSS", "Impostos", 12993.5, "Variável", "1", "STB", None, r"^inss\b"),
	("IRPS", "Impostos", 10280.75, "Variável", "1", "STB", None, r"^irps\b"),
	("IVA", "Impostos", 0, "Variável", "1", "STB", None, r"^iva\b"),
	("Salários", "Salários e Encargos", 221221.35, "Variável", "1", "STB", None, r"^salarios\b"),
	("Energia", "Instalações", 3000, "Fixo", "2", "STB", "Nr de contador 04240906836", r"^energia\b"),
	("Água - FIPAG", "Instalações", 3000, "Variável", "2", "BIM C", None, r"^agua\b"),
	("Renda", "Instalações", 25000, "Fixo", "2", "STB", "NIB: 000800004452990410195 (BCI)", r"^renda\b"),
	("AWS", "Cloud e TI", 98709.64, "Variável", "2", "STB", None, r"^aws\b"),
	("Starlink", "Cloud e TI", 3000, "Fixo", "2", "STB", None, r"^starlink\b"),
	(
		"Prestação de Serviços",
		"Serviços",
		5800,
		"Fixo",
		"2",
		"STB",
		"NIB: 000800008343652310113",
		r"^prestacao de servicos\b",
	),
	("Caixa", "Outros", 7500, "Fixo", "3", "BIM C", None, r"^caixa\b"),
	("Despesas com Deslocações", "Transporte e Combustível", 5000, "Fixo", "3", "BIM C", None, r"^despesas com desloc"),
	("Combustível", "Transporte e Combustível", 20000, "Fixo", "3", "BIM C", None, r"^combustivel\b"),
]

# Categories for one-off lines, first match wins.
CATEGORIAS_POR_PALAVRA = [
	(r"\birpc\b", "Impostos"),
	(r"reembolso", "Reembolsos"),
	(r"^investimento", "Investimento"),
	(r"aniversario|saida com a equipe|brindes|1 de maio|7 de abril|bonificacao|ajuda|camisetas|logotipo", "Pessoal e Eventos"),
	(r"seguro acidente", "Salários e Encargos"),
	(r"reseller|dominio|internet|\bnic\b", "Cloud e TI"),
	(r"^dw\b|agua", "Instalações"),
	(r"honorarios|juridic|seguro|mao de obra|manuntecao|manutencao|montagem|recondicionamento", "Serviços"),
	(
		r"computador|material|tonner|toner|\bac\b|geladeira|agendas|monitores|celular|cadeiras|secretarias|candeeiro|arranque",
		"Escritório e Equipamento",
	),
]

PADRAO_ORDEM_COMPRA = re.compile(r"PUR-ORD-\d{4}-\d+")


def importar(caminho, ano=2026, substituir=False):
	create_categorias_de_despesa()
	_criar_metodos_de_pagamento()
	_criar_despesas_recorrentes()

	resultado = []
	for plano in ler_cashflow(caminho, ano):
		nome = nome_plano(plano["tipo"], ano, plano.get("mes"))
		if frappe.db.exists("Plano de Pagamentos", nome):
			if not substituir:
				resultado.append(f"{nome}: já existe, ignorado")
				continue
			frappe.delete_doc("Plano de Pagamentos", nome, ignore_permissions=True)

		doc = frappe.get_doc(
			{
				"doctype": "Plano de Pagamentos",
				"tipo": plano["tipo"],
				"ano": ano,
				"mes": plano.get("mes"),
				"estado": _estado_do_plano(plano, ano),
				"linhas": plano["linhas"],
			}
		).insert(ignore_permissions=True)
		resultado.append(
			f"{doc.name}: {len(doc.linhas)} linhas · previsto {doc.total_previsto:,.2f}"
			f" · pago {doc.total_pago:,.2f} · remanescente {doc.total_remanescente:,.2f}"
		)
		resultado.extend(f"   ! {aviso}" for aviso in plano["avisos"])

	frappe.db.commit()
	print("\n".join(resultado))
	return resultado


# ----------------------------------------------------------------------
# Parsing (no database access — testable on its own)
# ----------------------------------------------------------------------


def ler_cashflow(caminho, ano):
	from openpyxl import load_workbook

	wb = load_workbook(caminho, data_only=True)
	planos = []
	for ws in wb.worksheets:
		titulo = _normalizar(ws.title)
		if titulo == "investimentos":
			planos.append({"tipo": "Investimentos", **_ler_folha(ws, investimento=True)})
			continue
		mes = next((m for m in MESES if titulo.startswith(_normalizar(m))), None)
		if mes:
			planos.append({"tipo": "Mensal", "mes": mes, **_ler_folha(ws)})
	return planos


def _ler_folha(ws, investimento=False):
	colunas = {}
	for cell in ws[3]:
		cabecalho = _normalizar(cell.value)
		if cabecalho.startswith("metodo"):
			colunas["metodo"] = cell.column
		elif cabecalho.startswith("observa"):
			colunas["observacoes"] = cell.column
	# Janeiro–Junho have a Factura column between Valor Pago and Método.
	colunas["factura"] = 6 if colunas.get("metodo") == 7 else None

	linhas, avisos = [], []
	for r in range(4, ws.max_row + 1):
		valor_celula = lambda col: ws.cell(r, col).value if col else None  # noqa: E731
		descricao = str(valor_celula(1) or "").strip()
		if _normalizar(descricao).startswith("total"):
			break
		if _normalizar(descricao) in ("", "0"):
			continue

		valor, aviso = _numero(valor_celula(2))
		if aviso:
			avisos.append(f"{descricao}: valor '{aviso}' estava como texto → {valor:,.2f}")
		valor_pago, _ = _numero(valor_celula(5))
		estado = ESTADOS.get(_normalizar(valor_celula(4)), "Pendente")

		if estado == "Pago" and not valor_pago:
			avisos.append(f"{descricao}: 'Pago' sem Valor Pago → assumido {valor:,.2f}")
		elif estado == "Pago" and valor_pago < valor:
			avisos.append(f"{descricao}: 'Pago' mas Valor Pago {valor_pago:,.2f} menor que o Valor {valor:,.2f} — verificar")
		elif valor_pago > valor:
			avisos.append(f"{descricao}: Valor Pago {valor_pago:,.2f} maior que o Valor {valor:,.2f} — verificar")

		observacoes = str(valor_celula(colunas.get("observacoes")) or "").strip()
		metodo = str(valor_celula(colunas.get("metodo")) or "").strip()
		prioridade = valor_celula(3)
		ordem_compra = PADRAO_ORDEM_COMPRA.search(observacoes)
		despesa_recorrente, categoria = _classificar(descricao)

		linhas.append(
			{
				"descricao": descricao[:140],
				"valor": valor,
				"valor_pago": valor_pago,
				"estado": estado,
				"prioridade": str(prioridade) if prioridade in (1, 2, 3, 4) else "",
				"metodo_pagamento": metodo if metodo in METODOS_DE_PAGAMENTO else None,
				"factura": str(valor_celula(colunas["factura"]) or "").strip() or None,
				"ordem_compra": ordem_compra.group(0) if ordem_compra else None,
				"observacoes": observacoes or None,
				"despesa_recorrente": None if investimento else despesa_recorrente,
				"categoria": "Investimento" if investimento else categoria,
			}
		)
	return {"linhas": linhas, "avisos": avisos}


def _classificar(descricao):
	texto = _normalizar(descricao)
	for nome, categoria, *_resto, padrao in DESPESAS_RECORRENTES:
		if re.search(padrao, texto):
			return nome, categoria
	for padrao, categoria in CATEGORIAS_POR_PALAVRA:
		if re.search(padrao, texto):
			return None, categoria
	return None, "Outros"


def _numero(valor):
	"""Returns (number, original_text_if_it_was_text). Handles the
	'70.057,02' / '96 942,22' / '1000,00' entries typed as text."""
	if valor is None:
		return 0.0, None
	if isinstance(valor, (int, float)):
		return float(valor), None
	texto = str(valor).replace("\xa0", "").replace(" ", "").strip()
	if texto in ("", "-"):
		return 0.0, None
	if "," in texto:
		texto = texto.replace(".", "").replace(",", ".")
	try:
		return float(texto), str(valor)
	except ValueError:
		return 0.0, str(valor)


def _normalizar(valor):
	texto = unicodedata.normalize("NFKD", str(valor or "")).encode("ascii", "ignore").decode()
	return " ".join(texto.lower().split())


# ----------------------------------------------------------------------
# Master data
# ----------------------------------------------------------------------


def _estado_do_plano(plano, ano):
	if plano["tipo"] != "Mensal":
		return "Em Curso"
	hoje = getdate()
	atual = (hoje.year, hoje.month)
	periodo = (ano, numero_mes(plano["mes"]))
	if periodo < atual:
		return "Fechado"
	return "Em Curso" if periodo == atual else "Rascunho"


def _criar_metodos_de_pagamento():
	for metodo in METODOS_DE_PAGAMENTO:
		if not frappe.db.exists("Mode of Payment", metodo):
			frappe.get_doc(
				{"doctype": "Mode of Payment", "mode_of_payment": metodo, "type": "Bank", "enabled": 1}
			).insert(ignore_permissions=True)


def _criar_despesas_recorrentes():
	for nome, categoria, valor, tipo_valor, prioridade, metodo, referencia, _padrao in DESPESAS_RECORRENTES:
		if frappe.db.exists("Despesa Recorrente", nome):
			continue
		frappe.get_doc(
			{
				"doctype": "Despesa Recorrente",
				"despesa": nome,
				"categoria": categoria,
				"valor_padrao": valor,
				"tipo_valor": tipo_valor,
				"prioridade": prioridade,
				"metodo_pagamento": metodo,
				"referencia": referencia,
				"ativo": 1,
			}
		).insert(ignore_permissions=True)
